"""Export the approved advisor layout, with portrait overview and wide tables.

uv run --no-project --isolated --offline --with pypandoc-binary --with python-docx \
    python paper_layout_advisor_2026_10_01/export_word.py
"""

from copy import deepcopy
import json
from pathlib import Path
import re
import sys
import tempfile
from zipfile import ZipFile

import pypandoc
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FOLDER = Path(__file__).resolve().parent
sys.path.insert(0, str(FOLDER.parent))
from scripts.export_paper_layout_word import absolute_links, make_reference

SOURCE = FOLDER / "LAYOUT.md"
OUTPUT = FOLDER / "LAYOUT_GUI_THAY_2026_10_02.docx"


def native_underline(value):
    """Turn Markdown's HTML ins markers into editable Pandoc Underline nodes."""
    if isinstance(value, dict):
        return {key: native_underline(item) for key, item in value.items()}
    if not isinstance(value, list):
        return value
    result = []
    index = 0
    while index < len(value):
        node = value[index]
        if isinstance(node, dict) and node.get("t") == "RawInline" and node.get("c") == ["html", "<ins>"]:
            end = index + 1
            while end < len(value) and not (
                isinstance(value[end], dict) and value[end].get("t") == "RawInline"
                and value[end].get("c") == ["html", "</ins>"]
            ):
                end += 1
            if end == len(value):
                raise ValueError("Unclosed underline marker")
            result.append({"t": "Underline", "c": native_underline(value[index + 1:end])})
            index = end + 1
        else:
            result.append(native_underline(node))
            index += 1
    return result


def format_document(path):
    document = Document(path)
    # Section properties describe the preceding section. Insert a portrait
    # next-page break before experiment details; the final section is landscape.
    details = next(p for p in document.paragraphs if p.text == "9. Thông tin chi tiết thí nghiệm đã chạy")
    portrait = deepcopy(document.sections[0]._sectPr)
    kind = portrait.find(qn("w:type"))
    if kind is None:
        kind = OxmlElement("w:type")
        portrait.append(kind)
    kind.set(qn("w:val"), "nextPage")
    boundary = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    properties.append(portrait)
    boundary.append(properties)
    details._p.addprevious(boundary)
    landscape = document.sections[-1]
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = Cm(29.7), Cm(21)
    landscape.left_margin = landscape.right_margin = Cm(1.4)
    landscape.top_margin = landscape.bottom_margin = Cm(1.5)
    for section in document.sections:
        section.header.paragraphs[0].text = "Uncertainty-aware RAG | Dàn bài gửi thầy | 02/10/2026"
        for run in section.header.paragraphs[0].runs:
            run.font.size = Pt(8)
    for paragraph in document.paragraphs:
        if re.match(r"9\.[2-6]\. ", paragraph.text) and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.page_break_before = True

    for table in document.tables:
        table.style = "Table Grid"
        table.autofit = False
        if len(table.columns) == 9:
            widths = [6.5] + [2.55] * 8
        elif table.cell(0, 0).text == "Method":
            widths = [14.9, 6.0, 6.0]
        else:
            widths = [3.2, 8.2, 6.0]
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row_index, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                if row_index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8EFF5")
                    cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.font.bold = True
    document.core_properties.title = "Dàn bài gửi thầy — Two-Stage Conformal Context Selection for RAG"
    document.core_properties.subject = "Tổng quan và chi tiết thí nghiệm — 02/10/2026"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.save(path)


def markdown_tables(markdown):
    tables = []
    for block in re.findall(r"(?:^\|[^\n]*\|\n?)+", markdown, re.MULTILINE):
        lines = block.splitlines()
        assert re.match(r"^\|[-:| ]+\|$", lines[1])
        tables.append([[cell.strip() for cell in line.strip("|").split("|")] for i, line in enumerate(lines) if i != 1])
    return tables


def validate(markdown, path):
    document = Document(path)
    source_headings = re.findall(r"^#{1,6}\s+(.+)$", markdown, re.MULTILINE)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == source_headings, "Heading order/content differs"
    tables = markdown_tables(markdown)
    assert len(document.tables) == len(tables) == 6
    bold_cells = underlined_cells = 0
    for table, expected in zip(document.tables, tables):
        assert len(table.rows) == len(expected)
        for row, source_cells in zip(table.rows, expected):
            for cell, source_cell in zip(row.cells, source_cells):
                text = re.sub(r"\*\*|</?ins>", "", source_cell)
                assert cell.text == text, (cell.text, text)
                runs = [run for p in cell.paragraphs for run in p.runs]
                if "**" in source_cell:
                    assert any(run.bold for run in runs), text
                    bold_cells += 1
                if "<ins>" in source_cell:
                    assert any(run.underline for run in runs), text
                    underlined_cells += 1
    assert [len(table.rows) - 1 for table in document.tables[1:]] == [15] * 5
    assert len(document.sections) == 2
    assert document.sections[0].orientation == WD_ORIENT.PORTRAIT
    assert document.sections[1].orientation == WD_ORIENT.LANDSCAPE
    body = "\n".join(p.text for p in document.paragraphs)
    assert "Đang làm việc cùng anh Hưng" in body
    assert "Tác giả:" not in body and "Hung Le" not in body
    assert not document.core_properties.author
    with ZipFile(path) as archive:
        assert archive.testzip() is None
        xml = archive.read("word/document.xml").decode()
        assert "<w:hyperlink" in xml
        assert "<ins>" not in xml
    print(f"Verified {len(headings)} headings, 6 tables, 75 result rows, {bold_cells} best and {underlined_cells} second-best cells.")


def main():
    markdown = SOURCE.read_text()
    ast = json.loads(pypandoc.convert_text(absolute_links(markdown, FOLDER), "json", format="markdown+tex_math_dollars+pipe_tables"))
    ast = native_underline(ast)
    with tempfile.TemporaryDirectory(prefix="advisor-docx-") as temporary:
        reference = Path(temporary) / "reference.docx"
        make_reference(reference)
        pypandoc.convert_text(json.dumps(ast, ensure_ascii=False), "docx", format="json", outputfile=str(OUTPUT), extra_args=[f"--reference-doc={reference}", "--standalone"])
    format_document(OUTPUT)
    validate(markdown, OUTPUT)
    print(f"Created {OUTPUT.name}: {OUTPUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
