#!/usr/bin/env python3
"""Export the paper layout to an editable Word document.

Run without modifying the project's ML environment:
    uv run --no-project --isolated --with pypandoc-binary --with python-docx \
        python scripts/export_paper_layout_word.py

Pandoc converts Markdown math to native Office Math and keeps tables, lists,
heading styles, and links. This script adds an A4 reference style and verifies
that every source heading and table made it into the output.
"""

from __future__ import annotations

import argparse
import re
import tempfile
from pathlib import Path
from urllib.parse import quote
from zipfile import ZipFile

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
GITHUB_ROOT = (
    "https://github.com/NguyenKhanh2603/Uncertainty-Aware-Iterative-RAG/"
    "blob/code/post-retrieval-cosine-three-baselines-2026-10-01/"
)


def make_reference(path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.1
    for name, size in [("Title", 21), ("Heading 1", 17), ("Heading 2", 14),
                       ("Heading 3", 12), ("Heading 4", 11)]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string("183A56")
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
    for name in ["List Bullet", "List Bullet 2", "List Bullet 3", "List Paragraph"]:
        document.styles[name].paragraph_format.space_after = Pt(3)
    header = section.header.paragraphs[0]
    header.text = "Uncertainty-aware RAG | Paper layout | 01-10-2026"
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Page ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(path)


def absolute_links(markdown: str, source_dir: Path) -> str:
    """Make local repository links usable when the Word file is shared."""
    def replace(match: re.Match[str]) -> str:
        label, target = match.groups()
        if re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target) or target.startswith("#"):
            return match.group(0)
        local, separator, fragment = target.partition("#")
        resolved = (source_dir / local).resolve()
        try:
            relative = resolved.relative_to(ROOT).as_posix()
        except ValueError:
            return match.group(0)
        url = GITHUB_ROOT + quote(relative, safe="/")
        if separator:
            url += "#" + fragment
        return f"[{label}]({url})"
    return re.sub(r"\[([^\]\n]+)\]\(([^)\s]+)\)", replace, markdown)


def polish_tables(path: Path) -> None:
    document = Document(path)
    for table in document.tables:
        table.style = "Table Grid"
        table.autofit = True
        if table.cell(0, 0).text == "Method" and len(table.columns) in (3, 5):
            # Result tables need a wider method column than the numeric columns.
            table.autofit = False
            first_width = 6.6 if len(table.columns) == 5 else 9.4
            widths = [first_width] + [(17.4 - first_width) / (len(table.columns) - 1)] * (len(table.columns) - 1)
            for column, width in zip(table.columns, widths):
                column.width = Cm(width)
            for row in table.rows:
                for cell, width in zip(row.cells, widths):
                    cell.width = Cm(width)
        # Repeat the column header if a long table spans pages.
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        header_properties.append(repeat)
        for row_index, row in enumerate(table.rows):
            no_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                if row_index == 0:
                    shade = OxmlElement("w:shd")
                    shade.set(qn("w:fill"), "E8EFF5")
                    cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(8 if len(table.columns) >= 6 else 9)
                        if row_index == 0:
                            run.font.bold = True
    document.core_properties.title = "Uncertainty-Aware RAG — Paper Layout"
    document.core_properties.subject = "Layout, writing plan, and supplied results — 2026-10-01"
    document.save(path)


def validate(markdown: str, output: Path) -> None:
    source_headings = re.findall(r"^#{1,6}\s+(.+)$", markdown, re.MULTILINE)
    expected_tables = len(re.findall(r"^\|\s*:?-{3,}.*\|\s*$", markdown, re.MULTILINE))
    document = Document(output)
    # python-docx .text omits native Office Math; include its text nodes as well.
    word_headings = [
        "".join(node.text or "" for node in p._p.iter() if node.tag in {qn("w:t"), qn("m:t")})
        for p in document.paragraphs if p.style.name.startswith("Heading")
    ]
    assert word_headings == [h.replace("$", "") for h in source_headings], "Heading content/order was lost in conversion"
    assert len(document.tables) == expected_tables, "Table count differs from Markdown"
    assert sum(len(t.rows) - 1 for t in document.tables[-9:]) == 135, "Supplied result rows missing"
    with ZipFile(output) as archive:
        assert archive.testzip() is None, "Invalid DOCX archive"
        xml = archive.read("word/document.xml").decode("utf-8")
        assert "<m:oMath" in xml, "Native Office Math equations missing"
        assert "<w:hyperlink" in xml, "Hyperlinks missing"
    print(f"Validated {len(word_headings)} headings, {len(document.tables)} tables, editable math and links.")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "PAPER_LAYOUT_UNCERTAINTY_AWARE_RAG.md")
    parser.add_argument("--output", type=Path, default=ROOT / "PAPER_LAYOUT_UNCERTAINTY_AWARE_RAG_2026_10_01.docx")
    args = parser.parse_args()
    source = args.source.resolve()
    markdown = source.read_text(encoding="utf-8")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="rag-word-") as folder:
        reference = Path(folder) / "reference.docx"
        make_reference(reference)
        pypandoc.convert_text(
            absolute_links(markdown, source.parent),
            "docx", format="markdown+tex_math_dollars+pipe_tables",
            outputfile=str(args.output),
            extra_args=[f"--reference-doc={reference}", "--standalone"],
        )
    polish_tables(args.output)
    validate(markdown, args.output)
    print(f"Created {args.output} ({args.output.stat().st_size:,} bytes).")


if __name__ == "__main__":
    main()
